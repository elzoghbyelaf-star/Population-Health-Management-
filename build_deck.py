#!/usr/bin/env python3
"""
Amigo x Cleveland Diagnostics — partnership deck.
Theme matched pixel-for-pixel to Amigo's own 'Population Health Intelligence'
deck: warm cream paper, Didone serif display, terracotta + navy accents,
letter-spaced monospace eyebrows/labels, near-flat cards, navy closing slide.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- Brand palette (sampled from Amigo_PopulationHealthIntelligence.pdf) ----
CREAM   = RGBColor(0xF1, 0xEC, 0xE4)   # page background
CARD    = RGBColor(0xF7, 0xF3, 0xEC)   # card fill (a touch lighter/warmer)
BORDER  = RGBColor(0xE1, 0xDA, 0xCF)   # hairline card border
INK     = RGBColor(0x22, 0x1D, 0x17)   # headlines / near-black
BODY    = RGBColor(0x47, 0x42, 0x3B)   # body copy
MUTE    = RGBColor(0x8B, 0x84, 0x78)   # citations / footer
TERRA   = RGBColor(0xB1, 0x4A, 0x2B)   # primary accent (terracotta)
PEACH   = RGBColor(0xCF, 0x97, 0x83)   # accent on navy
NAVY    = RGBColor(0x20, 0x30, 0x4A)   # deep navy
NAVYTXT = RGBColor(0xEF, 0xEA, 0xE2)   # headline on navy
NAVYBOD = RGBColor(0xC3, 0xBE, 0xB6)   # body on navy

SERIF = "Playfair Display"   # Didone display serif
SANS  = "Inter"              # humanist sans body
MONO  = "Space Mono"         # monospace eyebrows / labels

EMU = 914400
SW, SH = 13.333, 7.5

prs = Presentation()
prs.slide_width  = Emu(int(SW * EMU))
prs.slide_height = Emu(int(SH * EMU))
BLANK = prs.slide_layouts[6]


# ---------------- low-level helpers ----------------
def _solid(shape, color):
    shape.fill.solid(); shape.fill.fore_color.rgb = color


def _no_line(shape):
    shape.line.fill.background()


def _line(shape, color, w=0.75, alpha=None):
    shape.line.color.rgb = color
    shape.line.width = Pt(w)
    if alpha is not None:
        ln = shape.line._get_or_add_ln()
        srgb = ln.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        srgb.append(srgb.makeelement(qn("a:alpha"), {"val": str(int(alpha * 1000))}))


def _shadow(shape, blur=0.05, dist=0.028, alpha=93):
    spPr = shape._element.spPr
    ef = spPr.makeelement(qn("a:effectLst"), {})
    sh = spPr.makeelement(qn("a:outerShdw"),
                          {"blurRad": str(int(blur * EMU)), "dist": str(int(dist * EMU)),
                           "dir": "5400000", "rotWithShape": "0"})
    clr = spPr.makeelement(qn("a:srgbClr"), {"val": "221D17"})
    clr.append(spPr.makeelement(qn("a:alpha"), {"val": str((100 - alpha) * 1000)}))
    sh.append(clr); ef.append(sh); spPr.append(ef)


def slide(bg=CREAM):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.shadow.inherit = False; _no_line(r); _solid(r, bg)
    return s


def rrect(s, x, y, w, h, radius_in=0.11, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = max(0.02, min(0.5, radius_in / min(w, h)))
        except Exception:
            pass
    return sp


def bar(s, x, y, w, h, color):
    b = rrect(s, x, y, w, h, shape=MSO_SHAPE.RECTANGLE)
    _no_line(b); _solid(b, color); return b


def ring(s, cx, cy, r, color, alpha):
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r),
                           Inches(2 * r), Inches(2 * r))
    o.shadow.inherit = False
    o.fill.background()
    _line(o, color, w=1.0, alpha=alpha)
    return o


def card(s, x, y, w, h, radius=0.13, shadow=True):
    c = rrect(s, x, y, w, h, radius_in=radius)
    _solid(c, CARD); _line(c, BORDER, 1.0)
    if shadow:
        _shadow(c)
    return c


def R(t, sz, col=BODY, bold=False, fnt=SANS, ital=False):
    return (t, sz, col, bold, fnt, ital)


def txt(s, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        space_after=4, line_spacing=1.0, spc=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space_after)
        p.space_before = Pt(0); p.line_spacing = line_spacing
        for (t, sz, col, bold, fnt, ital) in para:
            r = p.add_run(); r.text = t
            f = r.font
            f.size = Pt(sz); f.bold = bold; f.italic = ital
            f.color.rgb = col; f.name = fnt
            if spc is not None:
                r._r.get_or_add_rPr().set("spc", str(int(spc * 100)))
    return tb


def footer(s, dark=False):
    lc = NAVYBOD if dark else MUTE
    txt(s, 0.7, SH - 0.52, 4, 0.3, [[R("AMIGO", 9, lc, False, MONO)]], spc=1.5)
    txt(s, SW - 6.7, SH - 0.52, 6, 0.3,
        [[R("POPULATION HEALTH INTELLIGENCE", 9, lc, False, MONO)]],
        align=PP_ALIGN.RIGHT, spc=1.5)


def head(s, kicker, title_paras, dark=False, y=0.7, tsize=34):
    kcol = PEACH if dark else TERRA
    tcol = NAVYTXT if dark else INK
    txt(s, 0.72, y, 11.5, 0.35, [[R(kicker.upper(), 11.5, kcol, False, MONO)]], spc=2)
    tp = [[R(seg[0], tsize, tcol, True, SERIF)] if isinstance(seg, tuple) else
          [R(seg, tsize, tcol, True, SERIF)] for seg in title_paras]
    txt(s, 0.7, y + 0.34, 11.9, 0.55 * len(title_paras) + 0.5, tp,
        line_spacing=0.98, space_after=0)
    bar(s, 0.74, y + 0.4 + 0.62 * len(title_paras), 1.45, 0.045, TERRA)


# ============================================================
# SLIDE 1 — TITLE  (cream, concentric rings — mirrors Amigo p1)
# ============================================================
s = slide()
for rr, al in [(3.9, 26), (3.0, 34), (2.15, 42), (1.35, 50)]:
    ring(s, 12.4, 2.75, rr, TERRA, al)
txt(s, 0.75, 0.9, 10, 0.4,
    [[R("PARTNERSHIP OPPORTUNITY  ·  AMIGO × CLEVELAND DIAGNOSTICS", 12, MUTE, False, MONO)]], spc=2)
txt(s, 0.7, 1.75, 10.4, 3.0,
    [[R("Scale IsoPSA", 62, INK, True, SERIF)],
     [R("to standard of care.", 62, INK, True, SERIF)]],
    line_spacing=0.96, space_after=0)
txt(s, 0.73, 4.35, 10, 0.9, [[R("Predict. Prevent. Act.", 34, TERRA, True, SERIF, True)]])
txt(s, 0.75, 5.75, 9, 0.4, [[R("PREPARED FOR", 11, MUTE, False, MONO)]], spc=2)
txt(s, 0.72, 6.05, 9, 0.5, [[R("Cleveland Diagnostics, Inc.", 22, NAVY, False, SERIF)]])
txt(s, 0.7, SH - 0.52, 6, 0.3, [[R("AMIGO · CONFIDENTIAL", 9, MUTE, False, MONO)]], spc=1.5)


# ============================================================
# SLIDE 2 — CLEVELAND DIAGNOSTICS SUMMARY
# ============================================================
s = slide(); footer(s)
head(s, "The Company", ["Cleveland Diagnostics at a glance."])

stats = [("2025", "FDA PMA approval of IsoPSA (Dec 1)", TERRA, "FDA · 2025"),
         ("$75M+", "Growth capital raised, led by Novo Holdings", NAVY, "Business Wire · 2024"),
         ("~90%", "Sensitivity for high-grade cancer", TERRA, "Klein et al. · 2022"),
         ("0.78", "Validation AUC, n≈1,093 across 8 sites", NAVY, "Urologic Oncology · 2022")]
tw, gap, x0, ty = 2.92, 0.18, 0.72, 2.15
for i, (big, lab, col, cite) in enumerate(stats):
    x = x0 + i * (tw + gap)
    card(s, x, ty, tw, 1.55)
    txt(s, x + 0.22, ty + 0.16, tw - 0.4, 0.6, [[R(big, 34, col, True, SERIF)]])
    txt(s, x + 0.24, ty + 0.72, tw - 0.44, 0.55, [[R(lab, 10, BODY)]], line_spacing=1.0)
    txt(s, x + 0.24, ty + 1.24, tw - 0.44, 0.25, [[R(cite, 8, MUTE, False, MONO)]])

# two profile cards
card(s, 0.72, 3.95, 6.03, 2.75)
txt(s, 0.98, 4.18, 5.5, 0.3, [[R("WHO THEY ARE", 10.5, TERRA, False, MONO)]], spc=1.5)
prof = [("Precision-oncology biotech", "Cleveland, OH — rooted in predecessor AnalizaDx."),
        ("Platform: SIA / IsoClear", "Reads protein structure, not concentration, to find cancer isoforms."),
        ("Lead product: IsoPSA", "Stratifies high-grade prostate-cancer risk; aids the biopsy decision."),
        ("New CEO", "Michael Iskra (Jan 2026), to drive commercial scale-up.")]
yy = 4.55
for h, d in prof:
    txt(s, 0.98, yy, 5.6, 0.3, [[R(h, 12.5, INK, True, SERIF), R("  —  " + d, 9.5, BODY)]],
        line_spacing=1.0)
    yy += 0.52

card(s, 6.97, 3.95, 5.63, 2.75)
txt(s, 7.23, 4.18, 5.1, 0.3, [[R("EVIDENCE & POSITIONING", 10.5, TERRA, False, MONO)]], spc=1.5)
ev = ["Prospective, multicenter validation (Klein et al., 2022)",
      "Beat total PSA and % free PSA on AUC and specificity",
      "SUO 2025: accuracy shown with and without mpMRI",
      "Competes with mpMRI pathways, 4Kscore, PHI"]
yy = 4.55
for e in ev:
    txt(s, 7.23, yy, 5.2, 0.35, [[R("→  ", 11, TERRA, True), R(e, 10.5, BODY)]])
    yy += 0.42
txt(s, 7.23, 6.28, 5.2, 0.3,
    [[R("Single product today — SIA platform can extend to other cancers.",
        9.5, MUTE, False, SANS, True)]])


# ============================================================
# SLIDE 3 — PAIN POINTS
# ============================================================
s = slide(); footer(s)
head(s, "The Challenge", ["The approval is won.", "The scale-up isn't."])
txt(s, 0.72, 2.45, 11.6, 0.4,
    [[R("Turning a validated test into standard-of-care at national scale is now the binding constraint.",
        13, BODY, False, SANS, True)]])

pains = [
    ("Commercial scale-up", "From a niche lab-developed test to nationwide IVD adoption — a salesforce can't reach every urologist and PCP."),
    ("Patient identification", "Eligible men (50+, rising PSA) surface in primary care, where risk-stratification is inconsistent and IsoPSA is rarely top-of-mind."),
    ("Adoption & reimbursement", "Ordering behavior is hard to shift; payer coverage remains unclear and slows the funnel."),
    ("Real-world evidence gap", "Payers and guideline bodies want real-world outcomes; key data (SUO 2025) is conference-stage, not yet peer-reviewed."),
    ("Competitive pressure", "mpMRI pathways, 4Kscore and PHI compete for the same biopsy-decision moment — differentiation must be re-proven."),
    ("Single-product risk", "Revenue rests entirely on IsoPSA; pipeline expansion needs large longitudinal datasets and validation partners."),
]
cw, ch, gx, gy = 3.9, 1.72, 0.2, 0.22
x0, y0 = 0.72, 2.95
for i, (h, d) in enumerate(pains):
    r, c = divmod(i, 3)
    x = x0 + c * (cw + gx); y = y0 + r * (ch + gy)
    card(s, x, y, cw, ch)
    txt(s, x + 0.25, y + 0.2, 1.0, 0.3, [[R("%02d" % (i + 1), 12, TERRA, False, MONO)]], spc=1)
    txt(s, x + 0.25, y + 0.46, cw - 0.5, 0.4, [[R(h, 14.5, INK, True, SERIF)]])
    txt(s, x + 0.25, y + 0.9, cw - 0.5, 0.75, [[R(d, 9.7, BODY)]], line_spacing=1.03)


# ============================================================
# SLIDE 4 — PAIN -> AMIGO AGENT MAPPING
# ============================================================
s = slide(); footer(s)
head(s, "The Fit", ["Every gap maps to an agent."])

cols = [(0.72, 3.5, "PAIN POINT"), (4.42, 2.75, "AMIGO AGENT"), (7.4, 5.2, "HOW IT SOLVES IT")]
hy = 2.35
for x, w, t in cols:
    txt(s, x, hy, w, 0.3, [[R(t, 10.5, TERRA, False, MONO)]], spc=1.5)
bar(s, 0.72, hy + 0.36, 11.9, 0.02, BORDER)

rows = [
    ("Finding eligible patients", "Risk Stratification",
     "Predictive scoring flags men 50+ with rising PSA across the population — a ready, self-refilling demand funnel for IsoPSA."),
    ("Patients slip through primary care", "Observation",
     "24/7 surveillance detects newly elevated PSA labs in real time and prompts an IsoPSA recommendation at the point of care."),
    ("Adoption & ordering behavior", "Campaign Tracking",
     "Closed-loop screening campaigns with live dashboards drive physician ordering and patient follow-through."),
    ("Real-world evidence gap", "Patient Profile (HIE)",
     "360° longitudinal records generate real-world outcomes data for payers, guidelines and label expansion."),
    ("Payer value & pipeline", "Simulation",
     "Digital-twin engine models biopsy-avoidance ROI for payer dossiers and forecasts new SIA biomarker programs."),
]
ry = hy + 0.5
rh = 0.82
for i, (pn, ag, hw) in enumerate(rows):
    y = ry + i * rh
    txt(s, 0.72, y + 0.06, 3.5, rh - 0.1, [[R(pn, 13, INK, True, SERIF)]], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 4.42, y + 0.06, 2.85, rh - 0.1,
        [[R("▸ ", 12, TERRA, True), R(ag, 12.5, NAVY, False, SERIF)]], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 7.4, y + 0.06, 5.2, rh - 0.1, [[R(hw, 10.3, BODY)]],
        anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.03)
    if i < len(rows) - 1:
        bar(s, 0.72, y + rh - 0.02, 11.9, 0.012, BORDER)


# ============================================================
# SLIDE 5 — HIGH-VALUE USE CASES
# ============================================================
s = slide(); footer(s)
head(s, "The Opportunity", ["Where the value compounds."])

ucs = [
    ("Population-scale demand generation",
     "Screen millions of longitudinal records to surface every IsoPSA-eligible man and route them to ordering providers."),
    ("Reflex screening pathway",
     "Elevated PSA in primary care auto-triggers an IsoPSA reflex recommendation — closing the gap from lab result to biopsy decision."),
    ("Real-world evidence engine",
     "HIE-powered outcomes tracking builds the peer-review-grade RWE dataset payers and guideline committees require."),
    ("Payer value modeling",
     "Simulate biopsy-avoidance and cost savings per population to arm reimbursement negotiations with defensible ROI."),
    ("Guideline & literature intelligence",
     "The NLP engine keeps medical affairs ahead of evidence, competitors and guideline shifts in prostate diagnostics."),
    ("Pipeline acceleration beyond IsoPSA",
     "Reuse Amigo's cancer cohorts and longitudinal data to validate future SIA biomarkers — de-risking single-product reliance."),
]
cw, ch, gx, gy = 3.9, 1.72, 0.2, 0.22
x0, y0 = 0.72, 2.2
for i, (h, d) in enumerate(ucs):
    r, c = divmod(i, 3)
    x = x0 + c * (cw + gx); y = y0 + r * (ch + gy)
    card(s, x, y, cw, ch)
    txt(s, x + 0.25, y + 0.2, 1.0, 0.3, [[R("%02d" % (i + 1), 12, TERRA, False, MONO)]], spc=1)
    txt(s, x + 0.25, y + 0.46, cw - 0.5, 0.55, [[R(h, 13.5, INK, True, SERIF)]], line_spacing=0.96)
    txt(s, x + 0.25, y + 1.02, cw - 0.5, 0.65, [[R(d, 9.7, BODY)]], line_spacing=1.03)


# ============================================================
# SLIDE 6 — VALUE & ROADMAP  (mirrors Amigo "THE PATH" p19)
# ============================================================
s = slide(); footer(s)
head(s, "The Path", ["From first agent to national screening reach."])

phases = [
    ("PHASE 01", "Workshop", "Co-define eligible-patient criteria and map HIE / lab data foundations with the clinical team.", "WEEKS 1–2"),
    ("PHASE 02", "First agent live", "Risk-stratification model runs on real data — 70–80% pre-built, the rest calibrated to the population.", "~6 WEEKS"),
    ("PHASE 03", "Reflex & campaigns", "Elevated-PSA surveillance triggers IsoPSA prompts; adoption campaigns run with live dashboards.", "QUARTERS"),
    ("PHASE 04", "Evidence & expand", "RWE accrues for payers; simulation powers value dossiers; extend to new SIA biomarkers.", "ONGOING"),
]
pw, gx, x0, py = 2.92, 0.18, 0.72, 2.75
for i, (ph, h, d, tf) in enumerate(phases):
    x = x0 + i * (pw + gx)
    card(s, x, py, pw, 2.55)
    txt(s, x + 0.24, py + 0.22, pw - 0.4, 0.3, [[R(ph, 10.5, TERRA, False, MONO)]], spc=1)
    txt(s, x + 0.24, py + 0.56, pw - 0.4, 0.6, [[R(h, 16, NAVY, True, SERIF)]], line_spacing=0.95)
    txt(s, x + 0.24, py + 1.18, pw - 0.46, 1.05, [[R(d, 9.8, BODY)]], line_spacing=1.05)
    txt(s, x + 0.24, py + 2.18, pw - 0.4, 0.3, [[R(tf, 9, MUTE, False, MONO)]], spc=1)
    if i < 3:
        txt(s, x + pw - 0.02, py + 0.95, 0.32, 0.4, [[R("→", 17, TERRA, True)]], align=PP_ALIGN.CENTER)

bar(s, 0.72, 5.55, 11.9, 0.02, BORDER)
txt(s, 0.72, 5.72, 11.9, 0.6,
    [[R("Deployment in ~6 weeks", 13, INK, True, SERIF),
      R("  ·  no model reaches production without clinical sign-off. We ship the loop — find the patient, prompt the order, prove the outcome.",
        12, BODY)]], line_spacing=1.1)


# ============================================================
# SLIDE 7 — AMIGO SUMMARY  (Predict -> Prevent -> Act loop, p9 style)
# ============================================================
s = slide(); footer(s)
head(s, "The Partner", ["Amigo — population health intelligence."])

loop = [
    ("01 · PREDICT", "Predict", "Live surveillance and risk stratification across a whole population — down to district, cohort and individual."),
    ("02 · PREVENT", "Prevent", "Simulate interventions before you spend — quantify the cost of inaction and rank options by return."),
    ("03 · ACT", "Act", "Launch the campaign, then track forecast-versus-observed in real time — closed-loop, not a data lake."),
]
cw, gx, x0, py = 3.86, 0.16, 0.72, 2.35
for i, (lab, h, d) in enumerate(loop):
    x = x0 + i * (cw + gx)
    card(s, x, py, cw, 2.2)
    txt(s, x + 0.26, py + 0.22, cw - 0.4, 0.3, [[R(lab, 10.5, TERRA, False, MONO)]], spc=1.5)
    txt(s, x + 0.26, py + 0.55, cw - 0.4, 0.55, [[R(h, 22, NAVY, True, SERIF)]])
    txt(s, x + 0.26, py + 1.15, cw - 0.5, 0.95, [[R(d, 10.3, BODY)]], line_spacing=1.05)
    if i < 2:
        txt(s, x + cw - 0.02, py + 0.85, 0.3, 0.4, [[R("→", 17, TERRA, True)]], align=PP_ALIGN.CENTER)

txt(s, 0.72, 4.72, 11.9, 0.3,
    [[R("OUTCOMES FEED BACK INTO PREDICTION", 9.5, MUTE, False, MONO)]], spc=2, align=PP_ALIGN.CENTER)
bar(s, 0.72, 5.02, 11.9, 0.012, BORDER)

strip = [("WHO WE ARE", "A healthcare-AI infrastructure company — safe, reliable AI across the care spectrum. Not a chatbot vendor."),
         ("PROVEN AT SCALE", "22.2x ROI · 6M+ patient interactions · deploys in ~6 weeks · SOC 2 · HIPAA · GDPR."),
         ("BACKED BY", "General Catalyst · Madrona · United Healthcare / Optum Ventures.")]
cw2, gx2, x0 = 3.86, 0.16, 0.72
for i, (lab, d) in enumerate(strip):
    x = x0 + i * (cw2 + gx2)
    txt(s, x, 5.28, cw2, 0.3, [[R(lab, 10, TERRA, False, MONO)]], spc=1.5)
    txt(s, x, 5.58, cw2 - 0.15, 1.0, [[R(d, 10, BODY)]], line_spacing=1.08)


# ============================================================
# SLIDE 8 — THE INVITATION  (navy, mirrors Amigo p20)
# ============================================================
s = slide(bg=NAVY); footer(s, dark=True)
txt(s, 0.75, 0.9, 10, 0.35, [[R("THE INVITATION", 11.5, PEACH, False, MONO)]], spc=2)
txt(s, 0.7, 1.35, 11.5, 1.8,
    [[R("From a validated test", 40, NAVYTXT, True, SERIF)],
     [R("to a screening engine you own.", 40, NAVYTXT, True, SERIF)]],
    line_spacing=0.98, space_after=0)
bar(s, 0.74, 3.15, 1.45, 0.045, PEACH)
txt(s, 0.72, 3.5, 11.6, 0.9,
    [[R("The evidence is settled and IsoPSA is approved. The missing piece is execution — finding the patient, prompting the order, and proving the outcome, as one closed loop built on population data.",
        14, NAVYBOD)]], line_spacing=1.2)

inv = [("WHY AMIGO", "Purpose-built for population-scale cancer screening — prostate risk-stratification is an extension, not a new build."),
       ("WHAT'S NEXT", "A workshop to scope eligible-patient criteria, then a first agent live on real data in ~6 weeks."),
       ("CONTACT", "elzoghby.elaf@amigo.ai")]
cw, gx, x0, iy = 3.86, 0.16, 0.72, 4.75
for i, (lab, d) in enumerate(inv):
    x = x0 + i * (cw + gx)
    txt(s, x, iy, cw, 0.3, [[R(lab, 10, PEACH, False, MONO)]], spc=1.5)
    col = PEACH if lab == "CONTACT" else NAVYBOD
    fnt = MONO if lab == "CONTACT" else SANS
    txt(s, x, iy + 0.32, cw - 0.15, 1.1, [[R(d, 10.5 if lab != "CONTACT" else 11, col, False, fnt)]],
        line_spacing=1.1)

txt(s, 0.72, 6.35, 11, 0.5,
    [[R("Let's build the screening engine for prostate cancer.  ", 18, NAVYTXT, False, SERIF),
      R("amigo.ai", 18, PEACH, False, SERIF)]])
txt(s, 0.7, SH - 0.52, 6, 0.3, [[R("AMIGO · CONFIDENTIAL", 9, NAVYBOD, False, MONO)]], spc=1.5)


out = "Amigo_x_ClevelandDiagnostics.pptx"
prs.save(out)
print("Saved", out, "with", len(prs.slides._sldIdLst), "slides")
