#!/usr/bin/env python3
"""Approximate renderer: read the actual .pptx shapes -> SVG -> PNG (for visual QA)."""
import sys, html
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import fitz

SCALE = 96.0 / 914400  # EMU -> px at 96dpi
prs = Presentation(sys.argv[1] if len(sys.argv) > 1 else "Amigo_x_ClevelandDiagnostics.pptx")
W = int(prs.slide_width * SCALE); H = int(prs.slide_height * SCALE)

FONTMAP = {"Playfair Display": "Georgia, 'Times New Roman', serif",
           "Inter": "'Helvetica Neue', Arial, sans-serif",
           "Space Mono": "'Courier New', monospace"}


def emu(v): return v * SCALE


def solid_rgb(fill):
    try:
        if fill.type is not None and fill.fore_color and fill.fore_color.rgb is not None:
            return "#%s" % str(fill.fore_color.rgb)
    except Exception:
        pass
    return None


def line_props(sh):
    col, alpha, w = None, 1.0, 1.0
    try:
        ln = sh._element.spPr.find(qn("a:ln"))
        if ln is not None:
            if ln.find(qn("a:noFill")) is not None:
                return None, 1.0, 0
            sf = ln.find(qn("a:solidFill"))
            if sf is not None:
                sc = sf.find(qn("a:srgbClr"))
                if sc is not None:
                    col = "#" + sc.get("val")
                    a = sc.find(qn("a:alpha"))
                    if a is not None:
                        alpha = int(a.get("val")) / 100000
            wv = ln.get("w")
            if wv:
                w = int(wv) * SCALE
    except Exception:
        pass
    return col, alpha, w


def has_nofill(sh):
    f = sh._element.spPr.find(qn("a:solidFill"))
    nf = sh._element.spPr.find(qn("a:noFill"))
    return nf is not None and f is None


for idx, slide in enumerate(prs.slides, 1):
    out = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" '
           f'viewBox="0 0 {W} {H}">']
    for sh in slide.shapes:
        try:
            x, y, w, h = emu(sh.left), emu(sh.top), emu(sh.width), emu(sh.height)
        except Exception:
            continue
        st = sh.shape_type
        # geometry shapes
        geom = sh._element.spPr.find(qn("a:prstGeom")) if sh._element.spPr is not None else None
        prst = geom.get("prst") if geom is not None else None
        if prst in ("rect", "roundRect", "ellipse"):
            fill = solid_rgb(sh.fill)
            lc, la, lw = line_props(sh)
            attrs = ''
            if fill and not has_nofill(sh):
                attrs += f' fill="{fill}"'
            else:
                attrs += ' fill="none"'
            if lc and lw:
                attrs += f' stroke="{lc}" stroke-width="{lw:.1f}" stroke-opacity="{la:.2f}"'
            if prst == "ellipse":
                out.append(f'<ellipse cx="{x+w/2:.1f}" cy="{y+h/2:.1f}" rx="{w/2:.1f}" ry="{h/2:.1f}"{attrs}/>')
            else:
                rx = 8 if prst == "roundRect" else 0
                out.append(f'<rect x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" rx="{rx}"{attrs}/>')
        # text
        if sh.has_text_frame:
            tf = sh.text_frame
            paras = tf.paragraphs
            # measure block height
            def psize(p):
                s = max([(r.font.size.pt if r.font.size else 12) for r in p.runs], default=12)
                return s
            lineheights = [psize(p) * (float(p.line_spacing) if p.line_spacing else 1.0) * 1.25 for p in paras]
            total = sum(lineheights)
            anchor = tf.vertical_anchor
            if anchor == MSO_ANCHOR.MIDDLE:
                cy = y + h / 2 - total / 2
            elif anchor == MSO_ANCHOR.BOTTOM:
                cy = y + h - total
            else:
                cy = y
            for p, lh in zip(paras, lineheights):
                if not p.runs:
                    cy += lh; continue
                sz = psize(p)
                baseline = cy + sz * 1.0
                al = p.alignment
                if al == PP_ALIGN.CENTER:
                    tx, anc = x + w / 2, "middle"
                elif al == PP_ALIGN.RIGHT:
                    tx, anc = x + w, "end"
                else:
                    tx, anc = x, "start"
                spans = []
                for r in p.runs:
                    col = "#000000"
                    try:
                        if r.font.color and r.font.color.rgb is not None:
                            col = "#%s" % str(r.font.color.rgb)
                    except Exception:
                        pass
                    fam = FONTMAP.get(r.font.name, "Arial, sans-serif")
                    rsz = r.font.size.pt if r.font.size else 12
                    fw = "700" if r.font.bold else "400"
                    fs = "italic" if r.font.italic else "normal"
                    spc = ''
                    rpr = r._r.find(qn("a:rPr"))
                    if rpr is not None and rpr.get("spc"):
                        spc = f' letter-spacing="{int(rpr.get("spc"))/100*SCALE*96/72:.1f}"'
                    spans.append(f'<tspan font-family="{fam}" font-size="{rsz*SCALE*96/72:.1f}" '
                                 f'font-weight="{fw}" font-style="{fs}" fill="{col}"{spc}>'
                                 f'{html.escape(r.text)}</tspan>')
                out.append(f'<text x="{tx:.1f}" y="{baseline:.1f}" text-anchor="{anc}">{"".join(spans)}</text>')
                cy += lh
    out.append('</svg>')
    svg = "\n".join(out)
    d = fitz.open(stream=svg.encode(), filetype="svg")
    pix = d[0].get_pixmap(matrix=fitz.Matrix(1.4, 1.4))
    pix.save(f"/tmp/qa/slide{idx}.png")
    print("rendered slide", idx)
